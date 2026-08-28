"""alerts.xlsx data layer. Uses a file lock so the web app (MCP server) and
the background worker can safely share the same spreadsheet."""
import os
import uuid
from datetime import datetime, timezone

from filelock import FileLock
from openpyxl import Workbook, load_workbook

from app.config import settings

COLUMNS = [
    "id", "condition", "target_price", "email", "note",
    "status", "created_at", "triggered_at", "triggered_price",
]


def _lock_path() -> str:
    return settings.EXCEL_PATH + ".lock"


def _ensure_workbook() -> None:
    os.makedirs(os.path.dirname(settings.EXCEL_PATH), exist_ok=True)
    if not os.path.exists(settings.EXCEL_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "alerts"
        ws.append(COLUMNS)
        wb.save(settings.EXCEL_PATH)


def _row_to_dict(row: tuple) -> dict:
    return {col: row[i] for i, col in enumerate(COLUMNS)}


def create_alert(condition: str, target_price: float, email: str = None, note: str = None) -> dict:
    condition = (condition or "").lower().strip()
    if condition not in ("above", "below"):
        raise ValueError('condition must be "above" or "below"')

    dest_email = email or settings.ALERT_TO_EMAIL
    if not dest_email:
        raise ValueError(
            "No destination email given, and ALERT_TO_EMAIL is not set in .env"
        )

    alert = {
        "id": uuid.uuid4().hex[:8],
        "condition": condition,
        "target_price": float(target_price),
        "email": dest_email,
        "note": note or "",
        "status": "active",
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "triggered_at": "",
        "triggered_price": "",
    }

    with FileLock(_lock_path()):
        _ensure_workbook()
        wb = load_workbook(settings.EXCEL_PATH)
        ws = wb["alerts"]
        ws.append([alert[c] for c in COLUMNS])
        wb.save(settings.EXCEL_PATH)

    return alert


def list_alerts(status: str = "all") -> list:
    with FileLock(_lock_path()):
        _ensure_workbook()
        wb = load_workbook(settings.EXCEL_PATH)
        ws = wb["alerts"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

    alerts = [_row_to_dict(r) for r in rows if r and r[0] is not None]
    if status and status != "all":
        alerts = [a for a in alerts if a["status"] == status]
    return alerts


def delete_alert(alert_id: str) -> bool:
    with FileLock(_lock_path()):
        _ensure_workbook()
        wb = load_workbook(settings.EXCEL_PATH)
        ws = wb["alerts"]
        target_idx = None
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            if row[0].value == alert_id:
                target_idx = i
                break
        if target_idx is None:
            return False
        ws.delete_rows(target_idx)
        wb.save(settings.EXCEL_PATH)
    return True


def mark_triggered(alert_id: str, price: float) -> None:
    with FileLock(_lock_path()):
        _ensure_workbook()
        wb = load_workbook(settings.EXCEL_PATH)
        ws = wb["alerts"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == alert_id:
                row[5].value = "triggered"
                row[7].value = datetime.now(timezone.utc).isoformat(timespec="seconds")
                row[8].value = price
                break
        wb.save(settings.EXCEL_PATH)
